from fastapi import APIRouter, Depends, Query
from fastapi.responses import StreamingResponse
from sqlalchemy.orm import Session
from typing import Optional
from datetime import datetime
import io
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

from app.db.session import get_db
from app.schemas.assy import AssyOrderQuery
from app.crud import assy as assy_crud

router = APIRouter()

@router.get("/export")
async def export_assy_list(
    db: Session = Depends(get_db),
    doc_no: Optional[str] = None,
    item_code: Optional[str] = None,
    supplier: Optional[str] = None,
    package_type: Optional[str] = None,
    is_closed: Optional[int] = None,
    order_date_start: Optional[datetime] = None,
    order_date_end: Optional[datetime] = None,
):
    # 构建查询参数
    query_params = AssyOrderQuery(
        doc_no=doc_no,
        item_code=item_code,
        supplier=supplier,
        package_type=package_type,
        is_closed=is_closed,
        order_date_start=order_date_start,
        order_date_end=order_date_end
    )
    
    # 获取数据
    orders = assy_crud.get_assy_list(db, query_params)
    
    # 创建Excel工作簿
    wb = Workbook()
    ws = wb.active
    ws.title = "装配订单列表"
    
    # 定义表头
    headers = [
        "订单号", "物料编码", "封装形式", "打印批号", "业务数量", 
        "收货数量", "在制数量", "加工方式", "成测程序", "打线图号",
        "线材", "备注", "订单日期", "到货日期", "供应商", "订单状态"
    ]
    
    # 设置表头样式
    header_font = Font(name='微软雅黑', bold=True, color='FFFFFF')
    header_fill = PatternFill(start_color='366092', end_color='366092', fill_type='solid')
    header_alignment = Alignment(horizontal='center', vertical='center')
    border = Border(
        left=Side(style='thin'),
        right=Side(style='thin'),
        top=Side(style='thin'),
        bottom=Side(style='thin')
    )
    
    # 写入表头
    for col, header in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col, value=header)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = header_alignment
        cell.border = border
    
    # 写入数据
    for row, order in enumerate(orders, 2):
        ws.cell(row=row, column=1, value=order.doc_no)
        ws.cell(row=row, column=2, value=order.item_code)
        ws.cell(row=row, column=3, value=order.z_package_type_name)
        ws.cell(row=row, column=4, value=order.lot_code)
        ws.cell(row=row, column=5, value=order.business_qty)
        ws.cell(row=row, column=6, value=order.receipted_price_qty)
        ws.cell(row=row, column=7, value=order.wip_qty)
        ws.cell(row=row, column=8, value=order.z_processing_purpose_name)
        ws.cell(row=row, column=9, value=order.z_testing_program_name)
        ws.cell(row=row, column=10, value=order.z_assembly_code)
        ws.cell(row=row, column=11, value=order.z_wire_name)
        ws.cell(row=row, column=12, value=order.remark)
        ws.cell(row=row, column=13, value=order.purchase_date)
        ws.cell(row=row, column=14, value=order.first_arrival_date)
        ws.cell(row=row, column=15, value=order.supplier_full_name)
        ws.cell(row=row, column=16, value="已结束" if order.receipt_close == 1 else "未结束")
        
        # 设置数据单元格样式
        for col in range(1, len(headers) + 1):
            cell = ws.cell(row=row, column=col)
            cell.alignment = Alignment(horizontal='center', vertical='center')
            cell.border = border
    
    # 调整列宽
    for col in range(1, len(headers) + 1):
        ws.column_dimensions[get_column_letter(col)].width = 15
    
    # 保存到内存
    excel_file = io.BytesIO()
    wb.save(excel_file)
    excel_file.seek(0)
    
    # 生成文件名
    current_time = datetime.now().strftime("%Y%m%d_%H%M%S")
    filename = f"装配订单列表_{current_time}.xlsx"
    
    # 返回文件流
    return StreamingResponse(
        excel_file,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={
            "Content-Disposition": f"attachment; filename={filename}"
        }
    ) 